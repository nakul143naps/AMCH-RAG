"""Layout-aware document loaders for TXT, Markdown, PDF, DOCX, PPTX, CSV, and HTML/Web."""

import csv
import logging
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any, Literal

import docx
import pypdf

from app.ingestion.tables import detect_and_format_text_table, format_table_as_markdown

logger = logging.getLogger(__name__)


class RawDocumentPart:
    """Represents a discrete part of an extracted document (e.g. text paragraph, table, chart)."""

    def __init__(
        self,
        content: str,
        page: int | None = None,
        section: str | None = None,
        modality: Literal["text", "table", "image_caption"] = "text",
        metadata: dict[str, Any] | None = None,
        image_bytes: bytes | None = None,
    ):
        self.content = content.strip()
        self.page = page
        self.section = section
        self.modality: Literal["text", "table", "image_caption"] = modality
        self.metadata = metadata or {}
        self.image_bytes = image_bytes


class BaseDocumentLoader(ABC):
    """Abstract base class for format-specific document loaders."""

    @abstractmethod
    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        """Extract structured parts from the given document file."""


class TextLoader(BaseDocumentLoader):
    """Loader for plain text (.txt) files."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        path = Path(file_path)
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        if not text.strip():
            return []

        # Check if entire file is a text table
        table_md = detect_and_format_text_table(text)
        if table_md:
            return [RawDocumentPart(content=table_md, page=1, modality="table")]

        return [RawDocumentPart(content=text, page=1, modality="text")]


class MarkdownLoader(BaseDocumentLoader):
    """Loader for Markdown (.md) files, preserving section hierarchy and tables."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        path = Path(file_path)
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        if not text.strip():
            return []

        lines = text.splitlines()
        parts: list[RawDocumentPart] = []
        current_section: str | None = None
        current_lines: list[str] = []

        def _flush(lines_to_flush: list[str], section: str | None) -> None:
            if not lines_to_flush:
                return
            content = "\n".join(lines_to_flush).strip()
            if not content:
                return
            # Detect if this block is a Markdown table
            if content.startswith("|") and "| ---" in content:
                parts.append(RawDocumentPart(content=content, section=section, page=1, modality="table"))
            else:
                parts.append(RawDocumentPart(content=content, section=section, page=1, modality="text"))

        for line in lines:
            if line.startswith(("# ", "## ", "### ")):
                _flush(current_lines, current_section)
                current_lines = []
                current_section = line.lstrip("#").strip()
            current_lines.append(line)

        _flush(current_lines, current_section)
        return parts if parts else [RawDocumentPart(content=text, page=1, modality="text")]


class PDFLoader(BaseDocumentLoader):
    """Layout-aware loader for PDF documents extracting text, tables, and embedded images."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        path = Path(file_path)
        parts: list[RawDocumentPart] = []

        reader = pypdf.PdfReader(str(path))
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                # Check for tabular blocks inside page text
                detected_table = detect_and_format_text_table(text)
                if detected_table:
                    parts.append(RawDocumentPart(content=detected_table, page=idx, modality="table"))
                else:
                    parts.append(RawDocumentPart(content=text, page=idx, modality="text"))

            # Extract embedded images / figures from page
            try:
                for img_idx, img_file in enumerate(page.images, start=1):
                    parts.append(
                        RawDocumentPart(
                            content="",
                            page=idx,
                            section=f"Figure {img_idx}",
                            modality="image_caption",
                            image_bytes=img_file.data,
                            metadata={"image_name": img_file.name, "page": idx},
                        )
                    )
            except Exception as e:  # noqa: BLE001
                logger.debug(f"PDFLoader image extraction skipped for page {idx}: {e}")

        return parts


class DocxLoader(BaseDocumentLoader):
    """Loader for DOCX files extracting text paragraphs and serialized Markdown tables."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        path = Path(file_path)
        doc = docx.Document(str(path))
        parts: list[RawDocumentPart] = []

        # Paragraphs (body text)
        lines = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text)

        if lines:
            parts.append(RawDocumentPart(content="\n\n".join(lines), page=1, modality="text"))

        # Tables serialized to Markdown
        for t_idx, table in enumerate(doc.tables, start=1):
            grid = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                grid.append(cells)
            if grid:
                table_md = format_table_as_markdown(grid)
                parts.append(
                    RawDocumentPart(
                        content=table_md,
                        page=1,
                        section=f"Table {t_idx}",
                        modality="table",
                    )
                )

        return parts


class PptxLoader(BaseDocumentLoader):
    """Loader for PowerPoint (.pptx) presentations extracting slide text, tables, and notes."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        from pptx import Presentation

        path = Path(file_path)
        prs = Presentation(str(path))
        parts: list[RawDocumentPart] = []

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = None
            text_lines = []

            for shape in slide.shapes:
                # Shape title
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        if not slide_title:
                            slide_title = text.splitlines()[0]
                        text_lines.append(text)

                # Slide tables
                if shape.has_table:
                    grid = []
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        grid.append(cells)
                    if grid:
                        table_md = format_table_as_markdown(grid)
                        parts.append(
                            RawDocumentPart(
                                content=table_md,
                                page=slide_idx,
                                section=slide_title or f"Slide {slide_idx}",
                                modality="table",
                            )
                        )

            # Slide notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    text_lines.append(f"Slide Notes: {notes_text}")

            if text_lines:
                parts.append(
                    RawDocumentPart(
                        content="\n\n".join(text_lines),
                        page=slide_idx,
                        section=slide_title or f"Slide {slide_idx}",
                        modality="text",
                    )
                )

        return parts


class CSVLoader(BaseDocumentLoader):
    """Loader for CSV files serializing tabular records into Markdown table chunks."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        path = Path(file_path)
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            grid = [row for row in reader if any(cell.strip() for cell in row)]

        if not grid:
            return []

        table_md = format_table_as_markdown(grid)
        return [RawDocumentPart(content=table_md, page=1, modality="table")]


class WebPageLoader(BaseDocumentLoader):
    """Loader for web pages and HTML files using trafilatura."""

    def load(self, file_path: str | Path) -> list[RawDocumentPart]:
        import trafilatura

        path = Path(file_path)
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            html_content = f.read()

        extracted = trafilatura.extract(
            html_content,
            include_tables=True,
            include_links=False,
            output_format="txt",
        )
        if not extracted or not extracted.strip():
            # Fallback to plain text extraction
            import re
            extracted = re.sub(r"<[^>]+>", " ", html_content)
            extracted = re.sub(r"\s+", " ", extracted).strip()

        if not extracted:
            return []

        return [RawDocumentPart(content=extracted, page=1, modality="text")]


def get_loader_for_file(file_path: str | Path) -> BaseDocumentLoader:
    """Factory to retrieve appropriate document loader by file extension."""
    ext = Path(file_path).suffix.lower()
    loaders: dict[str, BaseDocumentLoader] = {
        ".txt": TextLoader(),
        ".md": MarkdownLoader(),
        ".pdf": PDFLoader(),
        ".docx": DocxLoader(),
        ".pptx": PptxLoader(),
        ".csv": CSVLoader(),
        ".html": WebPageLoader(),
        ".htm": WebPageLoader(),
    }
    loader = loaders.get(ext)
    if not loader:
        raise ValueError(
            f"Unsupported file extension: {ext}. Supported: {list(loaders.keys())}"
        )
    return loader
