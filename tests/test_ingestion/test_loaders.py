"""Tests for document loaders (TXT, MD, DOCX, PDF)."""

from pathlib import Path

import docx
import pypdf
import pytest

from app.ingestion.loaders import (
    DocxLoader,
    MarkdownLoader,
    PDFLoader,
    TextLoader,
    get_loader_for_file,
)


def test_text_loader(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("Hello world! This is a simple test file.", encoding="utf-8")

    loader = TextLoader()
    parts = loader.load(file_path)

    assert len(parts) == 1
    assert "Hello world!" in parts[0].content
    assert parts[0].page == 1


def test_markdown_loader(tmp_path: Path):
    file_path = tmp_path / "sample.md"
    content = """# Header One
This is section one content.

## Header Two
This is section two content.
"""
    file_path.write_text(content, encoding="utf-8")

    loader = MarkdownLoader()
    parts = loader.load(file_path)

    assert len(parts) >= 2
    assert "section one" in parts[0].content
    assert parts[0].section == "Header One"
    assert "section two" in parts[1].content
    assert parts[1].section == "Header Two"


def test_docx_loader(tmp_path: Path):
    file_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_paragraph("First paragraph of docx document.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"
    table.cell(1, 1).text = "Active"
    doc.save(str(file_path))

    loader = DocxLoader()
    parts = loader.load(file_path)

    assert len(parts) == 2
    assert parts[0].modality == "text"
    assert "First paragraph" in parts[0].content
    assert parts[1].modality == "table"
    assert "Key | Value" in parts[1].content


def test_pdf_loader(tmp_path: Path):
    file_path = tmp_path / "sample.pdf"

    # Create a simple PDF using pypdf writer
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with open(file_path, "wb") as f:
        writer.write(f)

    loader = PDFLoader()
    # Blank page returns empty parts (which is expected for blank pages)
    parts = loader.load(file_path)
    assert isinstance(parts, list)


def test_get_loader_factory(tmp_path: Path):
    assert isinstance(get_loader_for_file("doc.txt"), TextLoader)
    assert isinstance(get_loader_for_file("doc.md"), MarkdownLoader)
    assert isinstance(get_loader_for_file("doc.pdf"), PDFLoader)
    assert isinstance(get_loader_for_file("doc.docx"), DocxLoader)

    with pytest.raises(ValueError):
        get_loader_for_file("doc.unsupported")
