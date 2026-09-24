"""Tests for Phase 9 Multi-Modal and Layout-Aware Ingestion."""

import io
from pathlib import Path
from unittest.mock import MagicMock, patch

from PIL import Image
from qdrant_client import QdrantClient

from app.ingestion.loaders import (
    CSVLoader,
    PDFLoader,
    PptxLoader,
    RawDocumentPart,
    WebPageLoader,
)
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.tables import detect_and_format_text_table, format_table_as_markdown
from app.ingestion.vision import VisionCaptioner
from app.retrieval.retriever import HybridRetriever
from app.retrieval.vector_store import VectorStoreManager


def test_table_serialization_helpers():
    # Grid markdown table formatting
    headers_and_rows = [
        ["Quarter", "Revenue ($M)", "Growth (%)"],
        ["Q1", "120.5", "+15%"],
        ["Q2", "145.2", "+20%"],
    ]
    md_table = format_table_as_markdown(headers_and_rows)
    assert "| Quarter | Revenue ($M) | Growth (%) |" in md_table
    assert "| --- | --- | --- |" in md_table
    assert "| Q1 | 120.5 | +15% |" in md_table
    assert "| Q2 | 145.2 | +20% |" in md_table

    # Unstructured text table detection and formatting
    pipe_text = "Metric | 2025 | 2026\nARR | $10M | $25M\nUsers | 50k | 120k"
    formatted = detect_and_format_text_table(pipe_text)
    assert formatted is not None
    assert "| Metric | 2025 | 2026 |" in formatted


def test_csv_loader(tmp_path: Path):
    csv_file = tmp_path / "financials.csv"
    csv_file.write_text(
        "Department,Budget,Spend,Variance\nEngineering,500000,480000,20000\nMarketing,300000,310000,-10000\n",
        encoding="utf-8",
    )

    loader = CSVLoader()
    parts = loader.load(csv_file)
    assert len(parts) == 1
    assert parts[0].modality == "table"
    assert "| Department | Budget | Spend | Variance |" in parts[0].content
    assert "| Engineering | 500000 | 480000 | 20000 |" in parts[0].content


def test_pptx_loader(tmp_path: Path):
    from pptx import Presentation
    from pptx.util import Inches

    pptx_file = tmp_path / "presentation.pptx"
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Add a table to the slide
    rows, cols = 2, 2
    left, top, width, height = Inches(1), Inches(1), Inches(4), Inches(2)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.cell(0, 0).text = "Phase"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Phase 9"
    table.cell(1, 1).text = "In Progress"

    prs.save(str(pptx_file))

    loader = PptxLoader()
    parts = loader.load(pptx_file)
    assert len(parts) >= 1
    table_part = next((p for p in parts if p.modality == "table"), None)
    assert table_part is not None
    assert "| Phase | Status |" in table_part.content
    assert "| Phase 9 | In Progress |" in table_part.content


def test_webpage_loader(tmp_path: Path):
    loader = WebPageLoader()
    html_content = """
    <html>
        <head><title>AMCH-RAG Architecture</title></head>
        <body>
            <article>
                <h1>Autonomous Agentic Architecture</h1>
                <p>The system leverages hybrid dense and sparse search combined with CRAG self-correction.</p>
                <table>
                    <tr><th>Component</th><th>Engine</th></tr>
                    <tr><td>Vector Search</td><td>Qdrant</td></tr>
                    <tr><td>Sparse Match</td><td>FastEmbed BM25</td></tr>
                </table>
            </article>
        </body>
    </html>
    """
    html_file = tmp_path / "article.html"
    html_file.write_text(html_content, encoding="utf-8")
    parts = loader.load(html_file)
    assert len(parts) >= 1
    content_combined = "\n".join(p.content for p in parts)
    assert "Autonomous Agentic Architecture" in content_combined or "Vector Search" in content_combined


def test_vision_captioner_gemini_integration():
    captioner = VisionCaptioner()

    # Create dummy 100x100 PNG
    img = Image.new("RGB", (100, 100), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()

    # Test fallback mode when GEMINI_API_KEY is empty
    with patch.object(captioner.settings, "GEMINI_API_KEY", ""):
        caption_fallback = captioner.caption_image_sync(
            img_bytes, mime_type="image/png", context_hint="Quarterly revenue chart"
        )
        assert "[Chart / Figure]" in caption_fallback
        assert "Quarterly revenue chart" in caption_fallback

    # Test online Gemini multimodal call with mock response
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.text = """[Figure: Bar Chart - Q4 Revenue Trends]
[Key Metrics: Q1: $10M, Q2: $15M, Q3: $25M, Q4: $40M]
[Trends: Exponential revenue growth observed over four quarters with peak in Q4]
[Caption: Bar chart illustrating rapid quarterly ARR acceleration from Q1 to Q4 peaking at $40M.]"""
    mock_client.models.generate_content.return_value = mock_resp

    with (
        patch.object(captioner.settings, "GEMINI_API_KEY", "fake-test-key"),
        patch("google.genai.Client", return_value=mock_client),
    ):
        caption = captioner.caption_image_sync(
            img_bytes, mime_type="image/png", context_hint="Financial performance report"
        )
        assert "Bar Chart - Q4 Revenue Trends" in caption
        assert "$40M" in caption


def test_multimodal_acceptance_pdf_ingestion_and_retrieval(tmp_path: Path):
    """Phase 9 Acceptance Criterion:

    A PDF containing a structured table and a chart figure is ingested such that questions
    about a table value and a chart trend are both answerable via hybrid retrieval.
    """
    test_client = QdrantClient(":memory:")
    vector_mgr = VectorStoreManager.get_instance(client=test_client)
    vector_mgr.ensure_collections()

    pipeline = IngestionPipeline(vector_mgr=vector_mgr)

    # 1. Create a synthetic image for a chart
    img = Image.new("RGB", (120, 120), color="red")
    img_buf = io.BytesIO()
    img.save(img_buf, format="PNG")
    chart_bytes = img_buf.getvalue()

    pdf_parts = [
        RawDocumentPart(
            content="Cloud Infrastructure Performance Report 2026. This report summarizes database latency and GPU utilization.",
            page=1,
            modality="text",
        ),
        RawDocumentPart(
            content="| Region | P99 Latency (ms) | Throughput (RPS) |\n| --- | --- | --- |\n| us-east-1 | 14.2 ms | 55,000 |\n| eu-west-1 | 18.7 ms | 38,000 |\n| ap-south-1 | 22.4 ms | 29,000 |",
            page=1,
            section="Table 1: Latency Benchmark",
            modality="table",
        ),
        RawDocumentPart(
            content="",
            page=1,
            section="Figure 1: GPU Utilization Curve",
            modality="image_caption",
            image_bytes=chart_bytes,
        ),
    ]

    # Mock VisionCaptioner to return detailed structured vision description
    mock_caption = """[Figure: Line Graph - GPU Utilization Curve]
[Key Metrics: Peak utilization 94% at 18:00 UTC, Baseline idle 12% at 04:00 UTC]
[Trends: Steady upward trajectory during trading hours from 08:00 to 18:00 UTC before plateauing]
[Caption: Line graph demonstrating GPU cluster compute saturation peaking at 94% during peak trading volume.]"""

    dummy_pdf = tmp_path / "system_benchmarks.pdf"
    dummy_pdf.write_bytes(b"%PDF-1.4 dummy content")

    with (
        patch.object(PDFLoader, "load", return_value=pdf_parts),
        patch.object(VisionCaptioner, "caption_image_sync", return_value=mock_caption),
    ):
        processed = pipeline.process_file(
            file_path=dummy_pdf,
            source_name="system_benchmarks.pdf",
            access_level="public",
        )

    assert processed.doc_id is not None
    assert len(processed.chunks) >= 3

    # Check that chunks in Qdrant maintain their modality metadata
    stored_chunks = vector_mgr.get_chunks_by_doc_id(processed.doc_id)
    modalities = {c.modality for c in stored_chunks}
    assert "text" in modalities
    assert "table" in modalities
    assert "image_caption" in modalities

    # Retrieve and verify table question answerability
    retriever = HybridRetriever(vector_mgr=vector_mgr)
    table_hits = retriever.retrieve(
        query="What is the P99 Latency in us-east-1?",
        limit=3,
        access_levels="public",
    )
    assert len(table_hits) > 0
    top_table_content = "\n".join(hit.content for hit in table_hits)
    assert "14.2 ms" in top_table_content
    assert "us-east-1" in top_table_content

    # Retrieve and verify chart trend question answerability
    chart_hits = retriever.retrieve(
        query="What is the peak GPU utilization trend and percentage?",
        limit=3,
        access_levels="public",
    )
    assert len(chart_hits) > 0
    top_chart_content = "\n".join(hit.content for hit in chart_hits)
    assert "94%" in top_chart_content or "GPU Utilization Curve" in top_chart_content
